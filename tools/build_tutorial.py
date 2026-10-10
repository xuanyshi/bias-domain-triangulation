"""Build the public HTML and Word tutorial from its Markdown source.
Requires python-docx. Run from any directory; does not alter analysis inputs.
"""
from pathlib import Path
import re, html, hashlib, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
STEM = 'Supplementary_Protocol_bias_domain_triangulation'
source = (DOCS / (STEM + '.md')).read_text()
blocks = []
lines = source.splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            if not re.match(r'^\|\s*-', lines[i]):
                rows.append([c.strip() for c in lines[i].strip('|').split('|')])
            i += 1
        blocks.append(('table', rows)); continue
    if line.startswith('#'):
        n = len(line) - len(line.lstrip('#'))
        blocks.append(('h' + str(n), line[n:].strip()))
    elif line.startswith('> '): blocks.append(('quote', line[2:]))
    elif line.startswith('- '): blocks.append(('ul', line[2:]))
    elif re.match(r'^\d+\. ', line): blocks.append(('ol', re.sub(r'^\d+\. ', '', line)))
    else: blocks.append(('p', line))
    i += 1

def inline(s):
    s = html.escape(s)
    s = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', s)
    s = re.sub(r'(?<!\*)\*([^*]+)\*', r'<em>\1</em>', s)
    return re.sub(r'https://[^\s<>]+', lambda m: '<a href="' + m[0] + '">' + m[0] + '</a>', s)

css = '''
:root{--ink:#202326;--muted:#5b626b;--line:#dbe1e5;--accent:#205982;--tint:#f1f6fa}
*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:24px}
body{margin:0;color:var(--ink);background:white;font:18px/1.65 "Times New Roman","STFangsong","华文仿宋",serif}
a{color:var(--accent);text-underline-offset:3px;overflow-wrap:anywhere}a:hover{color:#14344d}
.layout{display:grid;grid-template-columns:250px minmax(0,860px);max-width:1180px;margin:auto}
aside{position:sticky;top:0;align-self:start;height:100vh;overflow:auto;padding:32px 20px;border-right:1px solid var(--line)}
.brand{font-size:18px;font-weight:bold;margin:0 0 15px}.toc a{display:block;font-size:15px;line-height:1.35;text-decoration:none;padding:7px 0;color:var(--muted)}
.toc a.active{color:var(--accent);font-weight:bold}.toc a:focus-visible,a:focus-visible{outline:2px solid var(--accent);outline-offset:3px}
main{min-width:0;padding:40px 48px 80px}h1{font-size:36px;line-height:1.15;margin:12px 0 20px}h2{font-size:25px;line-height:1.3;margin:44px 0 16px;border-top:1px solid var(--line);padding-top:24px}p{margin:14px 0}li{margin:6px 0}
.kicker{color:var(--accent);font-size:14px;letter-spacing:1px;text-transform:uppercase}.release{padding:16px 20px;background:var(--tint);border-left:3px solid var(--accent);font-size:16px;margin:22px 0}.downloads{display:flex;flex-wrap:wrap;gap:8px 18px;margin-top:9px}
blockquote{background:#f7f8f9;border-left:3px solid var(--line);margin:20px 0;padding:10px 20px;font-size:17px}
.table-wrap{overflow-x:auto;margin:20px 0}table{width:100%;border-collapse:collapse;font-size:15px;line-height:1.4}th,td{text-align:left;vertical-align:top;padding:10px 9px;border-bottom:1px solid var(--line)}th{background:var(--tint);border-top:2px solid var(--accent)}td:first-child{font-weight:bold}footer{border-top:1px solid var(--line);margin-top:40px;padding-top:16px;font-size:15px;color:var(--muted)}
@media(max-width:850px){.layout{display:block}aside{position:static;height:auto;border-right:0;border-bottom:1px solid var(--line);padding:20px 22px}.toc{columns:2;column-gap:25px}.toc a{break-inside:avoid}.brand{margin-bottom:8px}main{padding:24px 22px 60px}h1{font-size:31px}h2{font-size:24px}.table-wrap{margin-right:0}table{min-width:520px}}
@media print{aside,.downloads{display:none}.layout{display:block}main{padding:0}body{font-size:11pt}h2{break-after:avoid}tr,blockquote{break-inside:avoid}a{color:inherit}.table-wrap{overflow:visible}}
'''
ids = ['overview'] + ['step-' + str(i) for i in range(1,11)] + ['reuse','references']
toc = []
body = []
hidx = 0
list_type = None
for kind, value in blocks:
    new_list = kind if kind in ('ul','ol') else None
    if new_list != list_type:
        if list_type: body.append('</' + list_type + '>')
        if new_list: body.append('<' + new_list + '>')
        list_type = new_list
    if kind == 'h2':
        anchor = ids[hidx]; hidx += 1
        toc.append('<a href="#' + anchor + '">' + html.escape(value) + '</a>')
        body.append('<h2 id="' + anchor + '">' + inline(value) + '</h2>')
    elif kind == 'h1': body.append('<h1>' + inline(value) + '</h1>')
    elif kind == 'table':
        rows = ['<thead><tr>' + ''.join('<th scope="col">'+inline(c)+'</th>' for c in value[0]) + '</tr></thead><tbody>']
        rows += ['<tr>'+''.join('<td>'+inline(c)+'</td>' for c in row)+'</tr>' for row in value[1:]]
        body.append('<div class="table-wrap" tabindex="0" role="region" aria-label="'+ ('Construct dictionary' if value[0][0]=='Node' else 'Grading criteria and examples') + '"><table>'+''.join(rows)+'</tbody></table></div>')
    elif kind in ('ul','ol'): body.append('<li>'+inline(value)+'</li>')
    elif kind == 'quote': body.append('<blockquote>'+inline(value)+'</blockquote>')
    else: body.append('<p>'+inline(value)+'</p>')
if list_type: body.append('</'+list_type+'>')
downloads = f'''<div class="release"><strong>Updated protocol · 10 October 2026</strong><br>Ten steps, two alternative adjustment sets, and explicit timing and interpretation boundaries.
<div class="downloads"><a href="{STEM}.docx">Word tutorial</a><a href="{STEM}.md">Markdown source</a><a href="assets/construct_dictionary.json">Construct dictionary</a><a href="assets/decision_rule.json">Grading rules</a><a href="assets/DAG132.dagitty">DAGitty graph</a></div></div>'''
page = '<!doctype html>\n<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="description" content="A ten-step tutorial for DAG-based construct coverage and bias-domain triangulation, updated 10 October 2026."><title>Bias-domain triangulation | Public tutorial</title><style>'+css+'</style></head><body><div class="layout"><aside><div class="brand">Bias-domain triangulation</div><nav class="toc" aria-label="Contents">'+''.join(toc)+'</nav><p style="font-size:14px"><a href="dag-review/">Earlier expert-review interface</a></p></aside><main><div class="kicker">Research methods · Reusable protocol</div>'+downloads+''.join(body)+'<footer><a href="https://github.com/xuanyshi/bias-domain-triangulation">Repository and version history</a> · <a href="assets/tutorial_manifest.json">Release manifest</a><p>The expert-review interface and older analysis files retain their own historical scope.</p></footer></main></div><script>const aliases={s0:"step-1",s1:"step-2",s2:"step-4","s2-2":"step-4","s2-3":"step-4",s3:"step-5",s4:"step-8",s5:"step-9",s6:"step-10",terms:"step-6"};const a=aliases[location.hash.slice(1)];if(a)location.replace("#"+a);const nav=[...document.querySelectorAll(".toc a")];const observer=new IntersectionObserver(es=>{for(const e of es)if(e.isIntersecting){nav.forEach(n=>n.classList.toggle("active",n.hash==="#"+e.target.id));}},{rootMargin:"-10% 0px -65% 0px"});document.querySelectorAll("h2[id]").forEach(x=>observer.observe(x));</script></body></html>'
(DOCS/'index.html').write_text(page)

# Word uses the same parsed source; font boundaries are explicit in OOXML.
d = Document()
for style in d.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec = d.sections[0]
sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.78)
for name in ['Normal','Title','Heading 1','Heading 2','List Bullet','List Number']:
    st = d.styles[name]
    st.font.name = 'Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'华文仿宋')
    st.font.size = Pt(11 if name not in ['Title','Heading 1','Heading 2'] else {'Title':24,'Heading 1':15,'Heading 2':12}[name])
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12
for name in ['Heading 1','Heading 2']:
    d.styles[name].paragraph_format.space_before=Pt(14)
    d.styles[name].paragraph_format.keep_with_next=True

def runs(p,s):
    for part in re.split(r'(\*\*.*?\*\*|\*[^*]+\*)',s):
        r=p.add_run(part.strip('*') if part.startswith('*') else part)
        r.bold=part.startswith('**'); r.italic=part.startswith('*') and not r.bold
        r.font.name='Times New Roman';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'华文仿宋')

number = 0
for kind,value in blocks:
    number = number + 1 if kind == 'ol' else 0
    if kind=='table':
        t=d.add_table(rows=1, cols=len(value[0]));t.autofit=False
        widths=([.45,3.5,2.76] if value[0][0]=='Node' else ([1.25,1.85,2.5,1.11] if len(value[0])==4 else [3.35,1.7,1.66]))
        for col,w in zip(t.columns,widths): col.width= Inches(w)
        for idx,row in enumerate(value):
            cells=t.rows[0].cells if idx==0 else t.add_row().cells
            for cell,w,text in zip(cells,widths,row):
                cell.width= Inches(w);p=cell.paragraphs[0];runs(p,text)
                p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
                for r in p.runs:r.font.size=Pt(9);r.bold=idx==0 or r.bold
                if idx==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F5');cell._tc.get_or_add_tcPr().append(sh)
            pr=t.rows[idx]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if idx==0:pr.append(OxmlElement('w:tblHeader'))
        d.add_paragraph()
    else:
        style={'h1':'Title','h2':'Heading 1','ul':'List Bullet'}.get(kind,'Normal')
        p=d.add_paragraph(style=style);runs(p,(str(number)+'. ' if kind=='ol' else '')+value)
        if kind=='ol':
            p.paragraph_format.left_indent= Inches(.2)
            p.paragraph_format.first_line_indent= Inches(-.2)
        if kind=='quote':p.paragraph_format.left_indent= Inches(.2)
d.core_properties.title='Bias-domain triangulation tutorial'
d.core_properties.subject='Public tutorial version 2026-10-10'
d.core_properties.author='';d.core_properties.last_modified_by=''
d.save(DOCS/(STEM+'.docx'))
manifest_path = DOCS/'assets/tutorial_manifest.json'
if manifest_path.exists():
    manifest = json.loads(manifest_path.read_text())
    manifest['files'] = {f:hashlib.sha256((DOCS/f).read_bytes()).hexdigest() for f in manifest['files']}
    manifest_path.write_text(json.dumps(manifest,indent=2)+'\n')
print('Built HTML and DOCX from',len(blocks),'source blocks')
